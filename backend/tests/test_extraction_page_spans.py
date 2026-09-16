from __future__ import annotations

import io
import unittest
from unittest.mock import patch

from docx import Document

from app.services.extraction import _extract_pdf_page_text, extract_text_with_pages


class _Page:
    def __init__(self, text: str | None) -> None:
        self._text = text

    def extract_text(self) -> str | None:
        return self._text


class _Pdf:
    def __init__(self, pages: list[_Page]) -> None:
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *_args) -> None:
        return None


class ExtractionPageSpanTests(unittest.TestCase):
    def test_pdf_character_layout_preserves_greek_super_and_subscripts(self) -> None:
        chars = [
            {"text": "q", "fontname": "Symbol", "size": 11, "x0": 0, "x1": 6, "top": 10},
            {"text": "v", "fontname": "Courier", "size": 11, "x0": 10, "x1": 16, "top": 10},
            {"text": "2", "fontname": "Courier", "size": 9, "x0": 16, "x1": 21, "top": 6},
            {"text": "q", "fontname": "Courier", "size": 11, "x0": 24, "x1": 30, "top": 10},
            {"text": "1", "fontname": "Courier", "size": 9, "x0": 30, "x1": 35, "top": 17},
        ]
        page = type("Page", (), {"chars": chars})()
        with patch(
            "app.services.extraction.extract_pdf_chars",
            side_effect=lambda mapped, **_kwargs: "".join(c["text"] for c in mapped),
        ) as extract:
            text = _extract_pdf_page_text(page)

        self.assertEqual(text, "θv²q₁")
        self.assertEqual(extract.call_args.kwargs["y_tolerance"], 8)

    def test_pdf_offsets_survive_outer_whitespace_trimming(self) -> None:
        fake_pdf = _Pdf([_Page("  page one"), _Page("page two  ")])
        with patch("app.services.extraction.pdfplumber.open", return_value=fake_pdf):
            result = extract_text_with_pages(b"pdf", "pdf")

        self.assertEqual(result.text, "page one\n\npage two")
        self.assertEqual(
            [(span.page_number, span.char_start, span.char_end) for span in result.page_spans],
            [(1, 0, 8), (2, 10, 18)],
        )

    def test_docx_keeps_existing_text_contract_without_fake_pages(self) -> None:
        with patch("app.services.extraction._extract_docx", return_value="body"):
            result = extract_text_with_pages(b"docx", "docx")
        self.assertEqual(result.text, "body")
        self.assertEqual(result.page_spans, ())

    def test_docx_extracts_paragraphs_and_table_rows_in_document_order(self) -> None:
        document = Document()
        document.add_paragraph("Before table")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Term"
        table.cell(0, 1).text = "Definition"
        table.cell(1, 0).text = "GDP"
        table.cell(1, 1).text = "Total output"
        document.add_paragraph("After table")
        payload = io.BytesIO()
        document.save(payload)

        result = extract_text_with_pages(payload.getvalue(), "docx")

        self.assertEqual(
            result.text,
            "Before table\nTerm | Definition\nGDP | Total output\nAfter table",
        )
        self.assertEqual(result.page_spans, ())


if __name__ == "__main__":
    unittest.main()
